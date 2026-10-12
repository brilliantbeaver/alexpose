#!/usr/bin/env python3
"""Render docs/faq.md as an illustrated Word document with SVG + PNG media.

This converter preserves the document's questions, headings, simple tables,
lists, source links and figure captions. SVG artwork is embedded through the
Office svgBlip extension; PNG fallbacks support viewers without SVG support.
"""
from __future__ import annotations

import argparse
import os
from pathlib import Path
import re
from urllib.parse import urlparse
from io import BytesIO
import zipfile

# Homebrew's shared libraries are outside the default macOS loader search.
if Path("/opt/homebrew/lib/libcairo.2.dylib").exists():
    os.environ.setdefault("DYLD_FALLBACK_LIBRARY_PATH", "/opt/homebrew/lib")
import cairosvg
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from lxml import etree
from markdown_it import MarkdownIt

ROOT = Path(__file__).resolve().parents[3]
SVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
SVG_EXT = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
REL_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
BLUE = "285A91"
INK = "183247"
MUTED = "526579"


def field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, code, separate, value, end):
        run._r.append(node)


def configure(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(.76)
    section.bottom_margin = Inches(.76)
    section.left_margin = Inches(.9)
    section.right_margin = Inches(.9)
    section.header_distance = Inches(.3)
    section.footer_distance = Inches(.3)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.widow_control = True
    for name, size, color in [("Title",27,INK),("Heading 1",21,INK),("Heading 2",16,BLUE),("Heading 3",12.3,INK)]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
    styles["Caption"].font.name = "Arial"
    styles["Caption"].font.size = Pt(9)
    styles["Caption"].font.color.rgb = RGBColor.from_string(MUTED)
    styles["Caption"].paragraph_format.keep_together = True
    styles["Caption"].paragraph_format.space_after = Pt(10)
    styles["Quote"].font.size = Pt(10)
    styles["Quote"].font.color.rgb = RGBColor.from_string(MUTED)
    styles["Quote"].paragraph_format.left_indent = Inches(.18)
    styles["Quote"].paragraph_format.right_indent = Inches(.1)
    for name in ["List Bullet", "List Number"]:
        styles[name].font.size = Pt(10.5)
        styles[name].paragraph_format.space_after = Pt(4)
    header = section.header.paragraphs[0]
    run=header.add_run("GAVD3  /  JEPA & GAIT RESEARCH FAQ")
    run.font.name="Arial"
    run.font.size=Pt(8)
    run.font.color.rgb=RGBColor.from_string(MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run=footer.add_run("Research concepts and evidence  •  ")
    run.font.size=Pt(8)
    run.font.color.rgb=RGBColor.from_string(MUTED)
    field(footer,"PAGE")
    document.core_properties.title="GAVD3: Frequently asked questions about JEPA and gait research"
    document.core_properties.subject="Concepts, methods, controls, findings and limits"
    document.core_properties.author="GAVD3 research project"
    document.core_properties.keywords="JEPA, S-JEPA, gait, representation learning, evidence audit"


def add_hyperlink(paragraph, text, target, bold=False, italic=False):
    link = OxmlElement("w:hyperlink")
    if target.startswith("#"):
        link.set(qn("w:anchor"), target[1:])
    else:
        rid = paragraph.part.relate_to(target,"http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",is_external=True)
        link.set(qn("r:id"),rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"),BLUE)
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"),"single")
    props.append(underline)
    if bold: props.append(OxmlElement("w:b"))
    if italic: props.append(OxmlElement("w:i"))
    run.append(props)
    txt = OxmlElement("w:t")
    txt.set(qn("xml:space"),"preserve")
    txt.text=text
    run.append(txt)
    link.append(run)
    paragraph._p.append(link)


def add_bookmark(paragraph, text, number):
    # Word bookmark names start with a letter, contain no hyphens and stay
    # within its 40-character limit. Source slugs map to these short names.
    name = f"q{number:03d}"
    start=OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"),str(number))
    start.set(qn("w:name"),name)
    end=OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"),str(number))
    paragraph._p.insert(0,start)
    paragraph._p.append(end)


class Renderer:
    def __init__(self, source, base_url=None):
        self.source=source
        self.base_url=base_url
        self.doc=Document()
        configure(self.doc)
        self.svg_count=0
        self.fig_count=0
        self.heading_count=0
        self.heading_anchors={}
        self.last_was_image=False

    def link_target(self, href):
        if href.startswith("#"):
            return "#"+self.heading_anchors.get(href[1:],href[1:])
        if urlparse(href).scheme:
            return href
        if self.base_url:
            relative=(self.source.parent / href).resolve().relative_to(ROOT)
            return self.base_url.rstrip("/")+"/"+str(relative)
        # Relative links resolve from the DOCX's docs/ directory, matching
        # Markdown. Keep the linked label readable, without appending paths.
        return href

    def inline(self, paragraph, children):
        bold=False; italic=False; target=None; link_text=""
        for token in children or []:
            typ=token.type
            if typ=="strong_open": bold=True
            elif typ=="strong_close": bold=False
            elif typ=="em_open": italic=True
            elif typ=="em_close": italic=False
            elif typ=="link_open":
                target=token.attrGet("href"); link_text=""
            elif typ=="link_close":
                resolved=self.link_target(target)
                add_hyperlink(paragraph,link_text,resolved,bold,italic)
                target=None
            elif typ=="image": self.image(token)
            elif typ in ("text","code_inline","softbreak","hardbreak"):
                content=" " if typ=="softbreak" else "\n" if typ=="hardbreak" else token.content
                if target is not None:
                    link_text+=content
                else:
                    run=paragraph.add_run(content)
                    # Preserve heading/caption style inheritance unless an
                    # inline marker explicitly adds emphasis.
                    run.bold=True if bold else None
                    run.italic=True if italic else None
                    if typ=="code_inline":
                        run.font.name="Consolas"
                        run.font.size=Pt(9)

    def image(self, token):
        path=(self.source.parent / token.attrGet("src")).resolve()
        if not path.is_file(): raise FileNotFoundError(path)
        alt=token.content or path.stem.replace("-"," ")
        p=self.doc.add_paragraph()
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(7)
        p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.keep_with_next=True
        if path.suffix.lower()==".svg":
            raw=path.read_bytes()
            png=cairosvg.svg2png(bytestring=raw,output_width=1800)
            shape=p.add_run().add_picture(BytesIO(png),width=Inches(6.4))
            self.svg_count+=1
            svgpart=Part(PackURI(f"/word/media/vector{self.svg_count}.svg"),"image/svg+xml",raw,self.doc.part.package)
            rid=self.doc.part.relate_to(svgpart,REL_IMAGE)
            blip=shape._inline.xpath(".//a:blip")[0]
            extlist=OxmlElement("a:extLst")
            ext=OxmlElement("a:ext")
            ext.set("uri",SVG_EXT)
            svgblip=etree.Element(f"{{{SVG_NS}}}svgBlip",nsmap={"asvg":SVG_NS})
            svgblip.set(qn("r:embed"),rid)
            ext.append(svgblip)
            extlist.append(ext)
            blip.append(extlist)
        else:
            shape=p.add_run().add_picture(str(path),width=Inches(6.4))
        shape._inline.docPr.set("descr",alt)
        shape._inline.docPr.set("title",path.stem.replace("-"," "))
        self.fig_count+=1
        self.last_was_image=True

    def table(self, tokens, start):
        rows=[]; row=[]; header=False; headers=[]; index=start+1
        while index<len(tokens) and tokens[index].type!="table_close":
            token=tokens[index]
            if token.type=="thead_open": header=True
            elif token.type=="thead_close": header=False
            elif token.type=="tr_open": row=[]
            elif token.type=="inline": row.append(token)
            elif token.type=="tr_close":
                rows.append(row)
                if header: headers.append(len(rows)-1)
            index+=1
        if rows:
            table=self.doc.add_table(rows=len(rows),cols=max(map(len,rows)))
            table.alignment=WD_TABLE_ALIGNMENT.CENTER
            table.style="Light Shading Accent 1"
            table.autofit=True
            for r,row in enumerate(rows):
                for c,token in enumerate(row):
                    cell=table.cell(r,c)
                    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p=cell.paragraphs[0]
                    self.inline(p,token.children)
                    p.paragraph_format.space_after=Pt(4)
                    p.paragraph_format.space_before=Pt(4)
                    for run in p.runs:
                        run.font.size=Pt(9)
                        if r in headers: run.bold=True
                trpr=table.rows[r]._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:cantSplit"))
                if r in headers: trpr.append(OxmlElement("w:tblHeader"))
            self.doc.add_paragraph().paragraph_format.space_after=Pt(1)
        return index

    def render(self):
        tokens=MarkdownIt("default",{"html":False}).parse(self.source.read_text(encoding="utf-8"))
        number=0
        for index,t in enumerate(tokens):
            if t.type=="heading_open":
                number+=1
                slug=re.sub(r"[^\w\- ]", "", tokens[index+1].content.lower()).replace(" ","-")
                self.heading_anchors[slug]=f"q{number:03d}"
        index=0; list_stack=[]; quoted=False
        while index<len(tokens):
            token=tokens[index]
            if token.type=="heading_open":
                level=int(token.tag[1])
                inline=tokens[index+1]
                style="Title" if level==1 and self.heading_count==0 else f"Heading {min(level,3)}"
                p=self.doc.add_paragraph(style=style)
                self.inline(p,inline.children)
                self.heading_count+=1
                add_bookmark(p,inline.content,self.heading_count)
                self.last_was_image=False
                index+=2
            elif token.type=="paragraph_open":
                inline=tokens[index+1]
                image_only=inline.children and all(t.type in ("image","softbreak") for t in inline.children)
                if image_only:
                    for image_token in inline.children:
                        if image_token.type=="image":self.image(image_token)
                else:
                    caption=self.last_was_image or inline.content.startswith("Figure ")
                    style="Caption" if caption else "Quote" if quoted else list_stack[-1] if list_stack else None
                    p=self.doc.add_paragraph(style=style)
                    self.inline(p,inline.children)
                    self.last_was_image=False
                index+=2
            elif token.type=="table_open": index=self.table(tokens,index)
            elif token.type in ("bullet_list_open","ordered_list_open"):
                list_stack.append("List Bullet" if token.type=="bullet_list_open" else "List Number")
            elif token.type in ("bullet_list_close","ordered_list_close"):
                list_stack.pop()
            elif token.type=="blockquote_open": quoted=True
            elif token.type=="blockquote_close": quoted=False
            elif token.type in ("fence","code_block"):
                p=self.doc.add_paragraph()
                run=p.add_run(token.content.rstrip())
                run.font.name="Consolas"
                run.font.size=Pt(9)
            elif token.type=="hr":
                p=self.doc.add_paragraph()
                p.paragraph_format.space_after=Pt(4)
            index+=1
        return self.doc


def validate(path, expected_vectors):
    ns={"a":"http://schemas.openxmlformats.org/drawingml/2006/main","asvg":SVG_NS,"r":"http://schemas.openxmlformats.org/officeDocument/2006/relationships","w":"http://schemas.openxmlformats.org/wordprocessingml/2006/main","wp":"http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}
    with zipfile.ZipFile(path) as archive:
        xml=etree.fromstring(archive.read("word/document.xml"))
        rels=etree.fromstring(archive.read("word/_rels/document.xml.rels"))
        relation_map={r.get("Id"):r.get("Target") for r in rels}
        vectors=xml.xpath("//asvg:svgBlip",namespaces=ns)
        assert len(vectors)==expected_vectors,(len(vectors),expected_vectors)
        for vector in vectors:
            target=relation_map[vector.get(qn("r:embed"))]
            assert "word/"+target in archive.namelist(),target
            assert target.endswith(".svg"),target
        for blip in xml.xpath("//a:blip",namespaces=ns):
            target=relation_map[blip.get(qn("r:embed"))]
            assert "word/"+target in archive.namelist(),target
            assert target.endswith(".png"),target
        drawings=xml.xpath("//wp:docPr",namespaces=ns)
        assert all(d.get("descr") for d in drawings),"A figure lacks alt text"
        assert "PAGE" in archive.read("word/footer1.xml").decode()
        bookmarks=xml.xpath("//w:bookmarkStart/@w:name",namespaces=ns)
        assert all(len(name)<=40 and re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*",name) for name in bookmarks)
        return {"svg_vectors":len(vectors),"png_fallbacks":len(drawings),"headings":len(xml.xpath("//w:pStyle[starts-with(@w:val,'Heading')]",namespaces=ns)),"tables":len(xml.xpath("//w:tbl",namespaces=ns))}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source",nargs="?",type=Path,default=ROOT/"docs/faq.md")
    parser.add_argument("--output",type=Path,default=ROOT/"docs/faq.docx")
    parser.add_argument("--base-url",help="Optional repository URL for relative source links")
    args=parser.parse_args()
    renderer=Renderer(args.source.resolve(),args.base_url)
    document=renderer.render()
    args.output.parent.mkdir(parents=True,exist_ok=True)
    document.save(args.output)
    report=validate(args.output,renderer.svg_count)
    print(f"Wrote {args.output}; validated {report}")


if __name__=="__main__": main()
