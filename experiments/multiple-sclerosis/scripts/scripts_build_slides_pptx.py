# /// script
# requires-python = ">=3.11"
# dependencies = ["python-pptx==1.0.2", "resvg-py==0.5.0", "beautifulsoup4==4.13.4"]
# ///
"""Build an editable PowerPoint from the September tutorial's Marp source.

Run from any directory:
    uv run --script scripts/scripts_build_slides_pptx.py

Slide titles, explanatory text, and tables remain editable. Illustrations retain
their original SVG data with a high-resolution PNG fallback for older viewers.
Marp speaker notes are copied into PowerPoint's native speaker-notes field.
The source Markdown and SVGs are the single source of scientific content.

The SVG fallback uses resvg-py, which bundles its renderer and needs no system
Cairo installation: https://resvg-py.readthedocs.io/en/latest/api.html
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
import hashlib
import html
from io import BytesIO
import json
import math
from pathlib import Path
import re
import zipfile

from lxml import etree
from PIL import ImageFont
from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
import resvg_py


ROOT = Path(__file__).resolve().parent.parent
NAVY = "142B45"
TEAL = "087E8B"
ORANGE = "C66C24"
MUTED = "536779"
PALE = "E6F2F3"
LINE = "CFDAE3"
WHITE = "FFFFFF"
FONT = "Arial"
WIDTH, HEIGHT = 13.333333, 7.5
LEFT, CONTENT_WIDTH = 0.70, 11.933333
SVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^\s)]+)(?:\s+[^)]*)?\)")
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<!\*)\*[^*]+\*(?!\*))")


def display_path(path: Path) -> str:
    return str(path.relative_to(ROOT)) if path.is_relative_to(ROOT) else str(path)


@dataclass
class SlideContent:
    title: str
    body: str
    images: list[tuple[str, Path]]
    notes: str
    kind: str = ""
    cards: list[tuple[str, str]] | None = None
    eyebrow: str = ""
    code_blocks: list[str] | None = None


def plain(value: str) -> str:
    value = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", value)
    value = re.sub(r"</?(?:span|div|small|strong|em|p)[^>]*>", "", value)
    value = re.sub(r"<br\s*/?>", " ", value)
    value = value.replace("**", "").replace("`", "")
    return html.unescape(value.strip())


def parse_source(source: Path) -> list[SlideContent]:
    text = source.read_text()
    if text.startswith("---\n"):
        text = text.split("\n---", 1)[1].lstrip("\n")
    slides = []
    for index, raw in enumerate(re.split(r"^---\s*$", text, flags=re.M), 1):
        if not raw.strip():
            continue
        comments = re.findall(r"<!--(.*?)-->", raw, flags=re.S)
        notes = []
        kind = ""
        for comment in comments:
            for line in comment.splitlines():
                if re.match(r"\s*_?class\s*:", line):
                    kind = line.split(":", 1)[1].strip()
            cleaned = "\n".join(
                line for line in comment.splitlines()
                if not re.match(r"\s*_?(?:class|paginate|header|footer|backgroundColor|color)\s*:", line)
            ).strip()
            if cleaned:
                notes.append(cleaned)
        visible = re.sub(r"<!--.*?-->", "", raw, flags=re.S)
        title_match = re.search(r"^#{1,2}\s+(.+)$", visible, flags=re.M)
        if not title_match:
            raise ValueError(f"Slide {index} has no level 1 or 2 heading")
        title = title_match.group(1).strip()
        visible = visible[:title_match.start()] + visible[title_match.end():]
        code_blocks = re.findall(r"```[^\n]*\n(.*?)```", visible, flags=re.S)
        visible = re.sub(r"```[^\n]*\n.*?```", "", visible, flags=re.S)
        soup = BeautifulSoup(visible, "html.parser")
        cards = []
        for group in soup.select(".cards"):
            for card in group.select(".card"):
                for br in card.find_all("br"):
                    br.replace_with("\n")
                heading = card.find(re.compile(r"^h[1-6]$"))
                card_title = heading.get_text(" ", strip=True) if heading else ""
                if heading:
                    heading.decompose()
                card_body = "\n".join(el.get_text().strip() for el in card.find_all(["p", "li"]))
                cards.append((card_title, card_body or card.get_text().strip()))
            group.decompose()
        eyebrow = ""
        for el in soup.select(".eyebrow"):
            eyebrow += el.get_text(" ", strip=True)
            el.decompose()
        for tag, delimiter in [("strong", "**"), ("b", "**"), ("em", "*"), ("code", "`")]:
            for el in soup.find_all(tag):
                el.replace_with(f"{delimiter}{el.get_text()}{delimiter}")
        for el in soup.find_all("a"):
            el.replace_with(f"[{el.get_text()}]({el.get('href', '')})")
        for el in soup.find_all("br"):
            el.replace_with(" ")
        # Inline HTML captions become ordinary editable text in PowerPoint.
        visible = str(soup)
        visible = re.sub(r"</(?:div|p|h[1-6])>", "\n\n", visible)
        visible = re.sub(r"<(?:div|p|h[1-6])[^>]*>", "", visible)
        images = [(alt, (source.parent / path).resolve()) for alt, path in IMAGE_RE.findall(visible)]
        visible = IMAGE_RE.sub("", visible)
        visible = re.sub(r"^#{1,6}\s+", "", visible, flags=re.M)
        visible = re.sub(r"</?(?:div|span|small|p)[^>]*>", "", visible)
        body = re.sub(r"\n{3,}", "\n\n", visible).strip()
        if re.search(r"<(?:img|table|svg)\b", body):
            raise ValueError(f"Slide {index}: use Markdown images/tables, not inline HTML")
        slides.append(SlideContent(title, body, images, "\n\n".join(notes), kind, cards, eyebrow, code_blocks))
    return slides


def font_file(bold: bool = False) -> str | None:
    names = [
        f"/System/Library/Fonts/Supplemental/Arial{' Bold' if bold else ''}.ttf",
        f"/usr/share/fonts/truetype/msttcorefonts/Arial{'_Bold' if bold else ''}.ttf",
        f"/usr/share/fonts/truetype/dejavu/DejaVuSans{'-Bold' if bold else ''}.ttf",
    ]
    return next((name for name in names if Path(name).exists()), None)


def line_count(text: str, size: float, width: float, bold: bool = False) -> int:
    path = font_file(bold)
    if not path:
        return sum(max(1, math.ceil(len(p) * size * 0.53 / (width * 72))) for p in text.splitlines())
    font = ImageFont.truetype(path, int(size * 4))
    limit = width * 72 * 4 * 0.98
    count = 0
    for paragraph in text.splitlines():
        line = ""
        count += 1
        for word in paragraph.split():
            candidate = f"{line} {word}".strip()
            if line and font.getlength(candidate) > limit:
                count += 1
                line = word
            else:
                line = candidate
    return max(count, 1)


def text_box(slide, value: str, x: float, y: float, w: float, h: float,
             size: float = 18, color: str = NAVY, bold: bool = False,
             align=PP_ALIGN.LEFT, name: str = "", valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = True
    frame.vertical_anchor = valign
    for i, line in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        p.font.name = FONT
        p.font.size = Pt(size)
        p.line_spacing = 1.13
        p.space_after = Pt(size * 0.34)
        line = re.sub(r"^[-*]\s+", "•  ", line)
        line = re.sub(r"^>\s*", "", line)
        for bit in INLINE_RE.split(line):
            if not bit:
                continue
            run = p.add_run()
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor.from_string(color)
            run.font.bold = bold or (bit.startswith("**") and bit.endswith("**"))
            if bit.startswith("`") and bit.endswith("`"):
                run.font.name = "Consolas"
                run.font.size = Pt(size * 0.94)
            if bit.startswith("*") and bit.endswith("*") and not bit.startswith("**"):
                run.font.italic = True
                bit = bit[1:-1]
            link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", bit)
            if link:
                run.text = html.unescape(link.group(1))
                if link.group(2).startswith(("https://", "http://")):
                    run.hyperlink.address = link.group(2)
                    run.font.color.rgb = RGBColor.from_string(TEAL)
                    run.font.underline = False
            else:
                run.text = plain(bit)
                # A run's adjacent space belongs to the source, not to plain().
                if bit.startswith(" "):
                    run.text = " " + run.text
                if bit.endswith(" "):
                    run.text += " "
    return shape


def rectangle(slide, x, y, w, h, fill, name=""):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.fill.background()
    shape._element.spPr.append(OxmlElement("a:effectLst"))
    for effect_ref in shape._element.xpath("./p:style/a:effectRef"):
        effect_ref.set("idx", "0")
    if name:
        shape.name = name
    return shape


def native_cards(slide, cards, x, y, w, h):
    columns = 2 if len(cards) == 4 else min(3, len(cards))
    rows = math.ceil(len(cards)/columns)
    gap = 0.22
    cw, ch = (w-gap*(columns-1))/columns, (h-gap*(rows-1))/rows
    for i, (heading, body) in enumerate(cards):
        xx, yy = x+(i%columns)*(cw+gap), y+(i//columns)*(ch+gap)
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(xx), Inches(yy), Inches(cw), Inches(ch))
        shape.name = f"Card background {i+1}"
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(PALE)
        shape.line.fill.background()
        shape._element.spPr.append(OxmlElement("a:effectLst"))
        for effect_ref in shape._element.xpath("./p:style/a:effectRef"):
            effect_ref.set("idx", "0")
        shape.adjustments[0] = 0.045
        heading_size = 23 if columns == 2 else 21
        heading_h = line_count(heading, heading_size, cw-0.44, True)*heading_size*1.13/72 + 0.08
        text_box(slide, heading, xx+0.22, yy+0.24, cw-0.44, heading_h, heading_size, NAVY, True, name=f"Editable card {i+1} title")
        body_size = 20 if columns == 2 else 18.5
        remaining = ch-heading_h-0.65
        while line_count(body, body_size, cw-0.44)*body_size*1.22/72 > remaining and body_size > 16:
            body_size -= 0.5
        text_box(slide, body, xx+0.22, yy+heading_h+0.40, cw-0.44, remaining, body_size, name=f"Editable card {i+1} body")


def native_code(slide, source, x, y, w):
    size = 14.5
    lines = source.strip("\n").splitlines()
    line_height = size*1.3
    h = len(lines)*line_height/72+0.32
    rectangle(slide, x, y, w, h, "EEF3F7", "Code background")
    shape = text_box(slide, "\n".join(lines), x+0.17, y+0.12, w-0.34, h-0.20, size, name="Editable reproduction commands")
    for p in shape.text_frame.paragraphs:
        p.space_after = 0
        p.line_spacing = Pt(line_height)
        p.font.name = "Consolas"
        p.font.size = Pt(size)
        for run in p.runs:
            run.font.name = "Consolas"
            if p.text.startswith("#"):
                run.font.color.rgb = RGBColor.from_string(MUTED)
    return h


def native_table(slide, lines: list[str], x, y, w, h):
    rows = [[v.strip() for v in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", v) for v in row)]
    if not rows:
        return
    ncols = len(rows[0])
    if any(len(row) != ncols for row in rows):
        raise ValueError("A Markdown table has inconsistent column counts")
    shape = slide.shapes.add_table(len(rows), ncols, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = "Editable data table"
    table = shape.table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.16)
            cell.margin_top = cell.margin_bottom = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(NAVY if r == 0 else PALE if r % 2 else WHITE)
            cell.text = plain(value)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.name = FONT
                    run.font.size = Pt(17)
                    run.font.bold = r == 0
                    run.font.color.rgb = RGBColor.from_string(WHITE if r == 0 else NAVY)
    return shape


def add_illustration(slide, path: Path, box, svg_parts, alt: str = ""):
    if not path.exists():
        raise FileNotFoundError(path)
    x, y, w, h = box
    if path.suffix.lower() != ".svg":
        from PIL import Image
        with Image.open(path) as image:
            ratio = image.width / image.height
        pw, ph = min(w, h * ratio), min(h, w / ratio)
        return slide.shapes.add_picture(str(path), Inches(x+(w-pw)/2), Inches(y+(h-ph)/2), width=Inches(pw), height=Inches(ph))
    raw = path.read_bytes()
    root = etree.fromstring(raw)
    viewbox = root.get("viewBox", "").split()
    iw, ih = map(float, viewbox[2:]) if len(viewbox) == 4 else (float(root.get("width")), float(root.get("height")))
    ratio = iw / ih
    pw, ph = min(w, h * ratio), min(h, w / ratio)
    # A 2,600-pixel fallback stays legible when opened by software without SVG support.
    png = resvg_py.svg_to_bytes(svg_string=raw.decode(), width=2600, font_family=FONT)
    picture = slide.shapes.add_picture(BytesIO(png), Inches(x+(w-pw)/2), Inches(y+(h-ph)/2), width=Inches(pw), height=Inches(ph))
    picture.name = f"Vector illustration: {path.stem}"
    title = root.find("{http://www.w3.org/2000/svg}title")
    desc = root.find("{http://www.w3.org/2000/svg}desc")
    description = " ".join(e.text for e in [title, desc] if e is not None and e.text)
    picture._element.nvPicPr.cNvPr.set("descr", description or alt or path.stem)
    key = hashlib.sha256(raw).hexdigest()
    package = slide.part.package
    if key not in svg_parts:
        svg_parts[key] = Part(package.next_partname("/ppt/media/vector%d.svg"), "image/svg+xml", package, raw)
    rid = slide.part.relate_to(svg_parts[key], RT.IMAGE)
    extension_list = OxmlElement("a:extLst")
    extension = OxmlElement("a:ext")
    extension.set("uri", "{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
    svg_blip = etree.SubElement(extension, f"{{{SVG_NS}}}svgBlip", nsmap={"asvg": SVG_NS})
    svg_blip.set(f"{{{REL_NS}}}embed", rid)
    extension_list.append(extension)
    picture._element.blipFill.blip.append(extension_list)
    return picture


def paragraphs(body: str) -> list[str]:
    """Respect list items; join editorial hard-wrapped prose into one paragraph."""
    result = []
    for block in re.split(r"\n\s*\n", body):
        if not block.strip():
            continue
        lines = block.strip().splitlines()
        if all(re.match(r"\s*(?:[-*]|\d+\.)\s+", line) for line in lines):
            result.extend(line.strip() for line in lines)
        else:
            result.append(" ".join(line.strip() for line in lines))
    return result


def canonical_text(value: str) -> str:
    value = re.sub(r"(?m)^\s*[-*>]\s+", "", plain(value))
    return " ".join(value.replace("*", "").replace("•", "").split())


def build(source: Path, output: Path) -> dict:
    source_bytes = source.read_bytes()
    contents = parse_source(source)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(WIDTH), Inches(HEIGHT)
    deck.core_properties.title = "Geometry, symmetry, and JEPA for human gait"
    deck.core_properties.subject = "September 13, 2026 tutorial and retained research results"
    deck.core_properties.author = "Gait representation study"
    deck.core_properties.keywords = "gait, symmetry, JEPA, source grouping, multiple sclerosis, Parkinson's"
    deck.core_properties.comments = "Built from slides.md; editable text, SVG illustrations, and native speaker notes."
    # Some viewers inherit hyperlink color and shape effects from the theme
    # despite direct formatting. Keep the theme consistent with the deck.
    drawing_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    for part in deck.part.package.iter_parts():
        if str(part.partname).startswith("/ppt/theme/"):
            theme = etree.fromstring(part.blob)
            for name, color in [("hlink", TEAL), ("folHlink", MUTED)]:
                node = theme.find(f".//{{{drawing_ns}}}clrScheme/{{{drawing_ns}}}{name}")
                if node is not None:
                    node.clear()
                    etree.SubElement(node, f"{{{drawing_ns}}}srgbClr").set("val", color)
            part._blob = etree.tostring(theme, xml_declaration=True, encoding="UTF-8", standalone=True)
    svg_parts = {}
    report = []
    for index, content in enumerate(contents, 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
        rectangle(slide, 0, 0, WIDTH, 0.065, NAVY, "Top rule")
        rectangle(slide, 0, 0, WIDTH * index / len(contents), 0.065, TEAL, "Tutorial progress")
        text_box(slide, content.eyebrow.upper() or "GAIT  /  GEOMETRY  /  JEPA", LEFT, 0.28, 9.3, 0.19, 9.5, TEAL, bold=True, name="Study label")
        text_box(slide, "13 SEP 2026", 10.6, 0.28, 2.03, 0.19, 9.5, MUTED, align=PP_ALIGN.RIGHT, name="Evidence date")
        title_size = 30 if index > 1 else 34
        while line_count(plain(content.title), title_size, CONTENT_WIDTH, True) > 2 and title_size > 25:
            title_size -= 1
        title_lines = line_count(plain(content.title), title_size, CONTENT_WIDTH, True)
        title_h = title_lines * title_size * 1.13 / 72 + 0.08
        text_box(slide, content.title, LEFT, 0.68, CONTENT_WIDTH, title_h, title_size, bold=True, name="Editable slide title")
        top = 0.68 + title_h + 0.20
        bottom = 6.96
        for code in content.code_blocks or []:
            top += native_code(slide, code, LEFT, top, CONTENT_WIDTH)+0.22
        table_lines = [line for line in content.body.splitlines() if line.strip().startswith("|")]
        body = "\n".join(line for line in content.body.splitlines() if not line.strip().startswith("|"))
        blocks = paragraphs(body)
        body_text = "\n".join(blocks)
        if content.cards:
            if content.images or table_lines:
                raise ValueError(f"Slide {index}: cards combined with illustrations/tables need an explicit layout")
            body_size = 18
            nlines = line_count(plain(body_text), body_size, CONTENT_WIDTH-0.17) if blocks else 0
            body_h = nlines*body_size*1.13/72 + max(0,len(blocks)-1)*body_size*0.34/72 + 0.18 if blocks else 0
            cards_h = min(4.3, bottom-top-body_h-0.24)
            native_cards(slide, content.cards, LEFT, top+0.06, CONTENT_WIDTH, cards_h)
            if blocks:
                text_top = bottom-body_h
                rectangle(slide, LEFT, text_top, 0.035, body_h-0.06, TEAL, "Explanation accent")
                text_box(slide, body_text, LEFT+0.17, text_top, CONTENT_WIDTH-0.17, body_h, body_size, name="Editable explanation")
        elif content.images:
            body_size = 17.5
            nlines = line_count(plain(body_text), body_size, CONTENT_WIDTH - 0.32) if blocks else 0
            body_h = nlines * body_size * 1.13 / 72 + max(0, len(blocks)-1) * body_size * 0.34 / 72 + 0.18 if blocks else 0
            while body_h > 1.65 and body_size > 15.5:
                body_size -= 0.5
                nlines = line_count(plain(body_text), body_size, CONTENT_WIDTH-0.32)
                body_h = nlines * body_size * 1.13 / 72 + max(0,len(blocks)-1)*body_size*0.34/72 + 0.18
            if body_h > 2.1:
                raise ValueError(f"Slide {index} has too much text below its illustration ({body_h:.2f} inches)")
            table_h = min(1.5, len(table_lines)*0.34) if table_lines else 0
            illustration_h = bottom - top - body_h - table_h - (0.18 if blocks else 0)
            if len(content.images) == 1:
                alt, path = content.images[0]
                add_illustration(slide, path, (LEFT, top, CONTENT_WIDTH, illustration_h), svg_parts, alt)
            else:
                for j, (alt, path) in enumerate(content.images):
                    slot = CONTENT_WIDTH / len(content.images)
                    add_illustration(slide, path, (LEFT+j*slot, top, slot-0.18, illustration_h), svg_parts, alt)
            text_top = bottom - body_h
            if table_lines:
                native_table(slide, table_lines, LEFT, text_top-table_h-0.05, CONTENT_WIDTH, table_h)
            if blocks:
                rectangle(slide, LEFT, text_top, 0.035, body_h - 0.06, TEAL, "Explanation accent")
                text_box(slide, body_text, LEFT+0.17, text_top, CONTENT_WIDTH-0.17, body_h, body_size, name="Editable explanation")
        else:
            body_size = 23 if len(body_text) < 500 else 20
            body_h = bottom - top - 0.1
            if table_lines:
                th = min(4.55, (len(table_lines)-1)*0.72)
                native_table(slide, table_lines, LEFT, top, CONTENT_WIDTH, th)
                top += th + 0.30
                body_h = bottom - top
            while line_count(plain(body_text), body_size, CONTENT_WIDTH-0.48)*body_size*1.25/72 > body_h and body_size > 16:
                body_size -= 0.5
            if blocks:
                # Keep separate citations on separate lines when a pair would
                # otherwise wrap a long hyperlink mid-word in some viewers.
                body_text = "\n".join(
                    re.sub(r";\s+(?=\[)", ";\n", block)
                    if len(re.findall(r"\]\(https?://",block))>1 and line_count(plain(block),body_size,CONTENT_WIDTH-0.48)>1
                    else block for block in blocks
                )
                rectangle(slide, LEFT, top+0.05, 0.04, max(0.5,body_h-0.3), TEAL, "Explanation accent")
                text_box(slide, body_text, LEFT+0.24, top+0.05, CONTENT_WIDTH-0.48, body_h-0.05, body_size, name="Editable explanation")
        rectangle(slide, LEFT, 7.13, CONTENT_WIDTH, 0.008, LINE, "Footer rule")
        text_box(slide, "Developmental study  •  retained results through September 13", LEFT, 7.23, 10.5, 0.17, 8.5, MUTED, name="Study scope footer")
        text_box(slide, f"{index:02d} / {len(contents):02d}", 11.6, 7.22, 1.03, 0.19, 9.5, MUTED, align=PP_ALIGN.RIGHT, name="Slide number")
        slide.notes_slide.notes_text_frame.text = content.notes
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left+shape.width > deck.slide_width+2 or shape.top+shape.height > deck.slide_height+2:
                raise ValueError(f"Slide {index} has a shape outside its bounds: {shape.name}")
        report.append({"slide":index, "title":plain(content.title), "illustrations":[{"path":display_path(p), "sha256":hashlib.sha256(p.read_bytes()).hexdigest()} for _, p in content.images], "cards":len(content.cards or []), "code_blocks":len(content.code_blocks or []), "notes_characters":len(content.notes), "body_font_points":body_size})
    output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(output)
    with zipfile.ZipFile(output) as archive:
        if archive.testzip() is not None:
            raise ValueError("PowerPoint ZIP checksum validation failed")
        names = archive.namelist()
        for name in names:
            if name.endswith((".xml", ".rels", ".svg")):
                etree.fromstring(archive.read(name))
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", name):
                slide_xml = etree.fromstring(archive.read(name))
                vector_blips = slide_xml.findall(f".//{{{SVG_NS}}}svgBlip")
                all_vector_blips = slide_xml.xpath("//*[local-name()='svgBlip']")
                if len(vector_blips) != len(all_vector_blips):
                    raise ValueError(f"An SVG picture uses an incorrect Office namespace in {name}")
                rel_name = str(Path(name).parent / "_rels" / (Path(name).name+".rels"))
                rel_xml = etree.fromstring(archive.read(rel_name))
                relations = {el.get("Id"):el for el in rel_xml}
                for blip in vector_blips:
                    rel = relations.get(blip.get(f"{{{REL_NS}}}embed"))
                    if rel is None or rel.get("Type") != RT.IMAGE or not rel.get("Target", "").endswith(".svg"):
                        raise ValueError(f"An SVG picture has no matching vector image relationship in {name}")
        svg_count = sum(name.endswith(".svg") for name in names)
        embedded_hashes = {hashlib.sha256(archive.read(name)).hexdigest() for name in names if name.endswith(".svg")}
        source_svg_hashes = {hashlib.sha256(path.read_bytes()).hexdigest() for content in contents for _,path in content.images if path.suffix.lower()==".svg"}
        if embedded_hashes != source_svg_hashes:
            raise ValueError("Embedded SVGs differ from the current source illustrations")
        note_count = sum(bool(re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)) for name in names)
    reopened = Presentation(output)
    if len(reopened.slides) != len(contents) or note_count != len(contents):
        raise ValueError("PowerPoint slide or speaker-note count differs from the source")
    for content, slide in zip(contents, reopened.slides):
        actual_notes = slide.notes_slide.notes_text_frame.text
        if actual_notes != content.notes:
            raise ValueError(f"Speaker notes were altered: {content.title}")
        actual = {shape.name:shape.text for shape in slide.shapes if shape.has_text_frame}
        expected = {"Editable slide title":content.title}
        body = "\n".join(line for line in content.body.splitlines() if not line.strip().startswith("|"))
        if body.strip():
            expected["Editable explanation"] = "\n".join(paragraphs(body))
        for i,(title,text) in enumerate(content.cards or [],1):
            expected[f"Editable card {i} title"] = title
            expected[f"Editable card {i} body"] = text
        for name,value in expected.items():
            if canonical_text(actual.get(name,"")) != canonical_text(value):
                raise ValueError(f"Source text changed on '{content.title}', in {name}")
        actual_code = [shape.text for shape in slide.shapes if shape.has_text_frame and shape.name == "Editable reproduction commands"]
        if actual_code != [code.strip("\n") for code in content.code_blocks or []]:
            raise ValueError(f"Reproduction commands changed on '{content.title}'")
        expected_cells = []
        for line in content.body.splitlines():
            if not line.strip().startswith("|"):
                continue
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?",cell) for cell in cells):
                expected_cells.extend(plain(cell) for cell in cells)
        actual_cells = [cell.text for shape in slide.shapes if shape.has_table for row in shape.table.rows for cell in row.cells]
        if actual_cells != expected_cells:
            raise ValueError(f"Editable table values changed on '{content.title}'")
    if source.read_bytes() != source_bytes:
        raise ValueError("The Markdown source changed during export; rebuild from the final source")
    return {"source":display_path(source), "source_sha256":hashlib.sha256(source_bytes).hexdigest(), "output":display_path(output), "output_sha256":hashlib.sha256(output.read_bytes()).hexdigest(), "slides":len(contents), "svg_assets":svg_count, "svg_namespace":SVG_NS, "speaker_notes":note_count, "text_preserved":True, "tables_and_commands_preserved":True, "source_svgs_preserved":True, "package_xml_checked":True, "shape_bounds_checked":True, "size_bytes":output.stat().st_size, "slide_details":report}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT/"slides/slides.md")
    parser.add_argument("--output", type=Path, default=ROOT/"slides/slides.pptx")
    parser.add_argument("--manifest", type=Path, help="Optional JSON record of content preservation and package validation")
    args = parser.parse_args()
    result = build(args.source.resolve(), args.output.resolve())
    if args.manifest:
        args.manifest.parent.mkdir(parents=True, exist_ok=True)
        args.manifest.write_text(json.dumps(result, indent=2)+"\n")
    print(json.dumps({k:v for k,v in result.items() if k != "slide_details"}, indent=2))


if __name__ == "__main__":
    main()
